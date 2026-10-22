"""
Automated test for the cargo-list loader.

For every CARGO row in the Excel sheet it works out, on its own, what the fare matrix should contain,
loads the sheet into an empty test database with `load_cargo_data`, and then walks through the screen
a user would use:

    1. pick the Campus      (dropdown)
    2. pick the City        (dropdown - must list the expected city)
    3. pick the Vehicle     (dropdown - must list the expected vehicle for that city)
    4. read the fare        (must equal the expected price)

Results are written to JSON and Excel in `test_reports/` (PASS / FAIL, and expected vs actual on failure).

Run (from the folder with manage.py):
    python manage.py test cashier.test

Settings (environment variables, all optional):
    CARGO_XLSX         path to the sheet          (default: cashier/cargo_list.xlsx)
    CARGO_REPORT_DIR   where reports are written  (default: cashier/test_reports)

The test database is created and destroyed by Django; your real data is never touched.
"""
import json
import os
from collections import Counter, defaultdict
from datetime import datetime
from decimal import Decimal
from io import StringIO
from pathlib import Path

import openpyxl
from django.apps import apps
from django.core.management import call_command
from django.test import TestCase

from .cargo_import import clean, norm, squash

APP = apps.get_containing_app_config(__name__).label
CAMPUS_ID = 1
CAMPUS_NAME = "Main"
HERE = Path(__file__).resolve().parent
XLSX = Path(os.environ.get("CARGO_XLSX", HERE / "cargo_list.xlsx"))
REPORT_DIR = Path(os.environ.get("CARGO_REPORT_DIR", HERE / "test_reports"))

# Sheet headers
H_TYPE, H_CAMPUS, H_DISTRICT, H_PICKUP = "BUS/CARGO/CAR POOL NO.", "CAMPUS", "DISTRICT", "PICK UP LOCATION"
H_VEHICLE, H_COST, H_REBATE = "Bus Company or Private Vehicle Type", "TOTAL COST PF RENT", "REBATE AMOUNT"

# The spec, written out again here so the test does not just re-run the loader's own code.
CAMPUS_TO_CITY_PREFIX = {"SRC": "Santa Rosa", "MAKATI": "Main"}
VEHICLE_ALIASES = {"SUVS": "SUV"}
PLACEHOLDER_VEHICLE = "Unidentified"
CENT = Decimal("0.01")


# ---------------------------------------------------------------- expected values, straight from the sheet
def _money(value):
    digits = "".join(ch for ch in str(value if value is not None else "") if ch.isdigit() or ch == ".")
    return Decimal(digits) if digits else None


def _expected_price(row):
    cost, rebate = _money(row.get(H_COST)), _money(row.get(H_REBATE))
    if cost:
        return cost.quantize(CENT), rebate == cost
    if rebate:
        return (rebate / Decimal("0.5")).quantize(CENT), False
    return Decimal("0.00"), False


def _expected_city(row):
    campus = clean(row.get(H_CAMPUS))
    prefix = CAMPUS_TO_CITY_PREFIX.get(squash(campus), campus)
    district = clean(row.get(H_DISTRICT))
    if not district or district.upper() == "DISTRICT":
        district = clean(row.get(H_PICKUP))
    return (f"{prefix} ({district})" if prefix and district else prefix or district) or None


def _expected_vehicles(row):
    names = []
    for part in clean(row.get(H_VEHICLE)).split("/"):
        part = part.strip()
        if part:
            part = VEHICLE_ALIASES.get(squash(part), part)
            if norm(part) not in [norm(n) for n in names]:
                names.append(part)
    return names or [PLACEHOLDER_VEHICLE]


def read_expected(path):
    """-> (groups, student_rows). groups: {(city key, vehicle key): {...}} one per city + vehicle."""
    ws = openpyxl.load_workbook(path, data_only=True).worksheets[0]
    it = ws.iter_rows(values_only=True)
    header = [clean(h) for h in next(it)]
    groups, student_rows = {}, []
    for excel_row, values in enumerate(it, start=2):
        row = dict(zip(header, values))
        if squash(row.get(H_TYPE)) != "CARGO":
            continue
        city = _expected_city(row)
        if not city:
            continue
        price, is_student = _expected_price(row)
        if is_student:
            student_rows.append(excel_row)
        for vehicle in _expected_vehicles(row):
            g = groups.setdefault(
                (norm(city), norm(vehicle)),
                {"city": city, "vehicle": vehicle, "rows": [], "prices": []},
            )
            g["rows"].append(excel_row)
            g["prices"].append(price)
    for g in groups.values():
        counts = Counter(p for p in g["prices"] if p) or Counter(g["prices"])  # 0.00 never beats a real price
        g["expected"] = max(counts, key=lambda p: (counts[p], p))
        g["conflict"] = dict(counts) if len(counts) > 1 else None
    return groups, student_rows


# ---------------------------------------------------------------- the dropdown walk-through
def simulate_dropdowns(Matrix, campus, city_key, vehicle_key):
    """Returns (failed_step | None, actual_price | None, detail)."""
    rows = list(Matrix.objects.filter(campus=campus).select_related("city", "vehicle"))
    # Step 1 - Campus dropdown
    if not rows:
        return 1, None, f"campus '{campus.name}' has no fares in the matrix"
    # Step 2 - City dropdown
    cities = {norm(m.city.name): m.city.name for m in rows}
    if city_key not in cities:
        return 2, None, f"city not in dropdown ({len(cities)} cities listed, none match)"
    # Step 3 - Vehicle dropdown for that city
    vehicles = {norm(m.vehicle.name): m for m in rows if norm(m.city.name) == city_key}
    if vehicle_key not in vehicles:
        return 3, None, f"vehicle not in dropdown; options: {sorted(m.vehicle.name for m in vehicles.values())}"
    # Step 4 - Fare
    return None, vehicles[vehicle_key].peso_price, ""


# ---------------------------------------------------------------- reports
def _result(test, parameters, expected, actual, passed, failed_step=None, remark=""):
    return {
        "test": test,
        "parameters": parameters,
        "expected": expected,
        "actual": actual,
        "result": "PASS" if passed else "FAIL",
        "failed_step": failed_step,
        "remark": remark,
    }


def write_reports(results):
    REPORT_DIR.mkdir(parents=True, exist_ok=True)
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    for i, r in enumerate(results, start=1):
        r["no"] = i
    passed = sum(r["result"] == "PASS" for r in results)
    summary = {
        "run_at": datetime.now().isoformat(timespec="seconds"),
        "excel_file": str(XLSX),
        "total": len(results), "passed": passed, "failed": len(results) - passed,
    }

    json_path = REPORT_DIR / f"cargo_test_report_{stamp}.json"
    json_path.write_text(json.dumps({"summary": summary, "results": results}, indent=2, ensure_ascii=False, default=str), encoding="utf-8")

    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Results"
    cols = ["No", "Test", "Campus", "City", "Vehicle", "Excel rows", "Expected", "Actual", "Result", "Failed step", "Remark"]
    ws.append(cols)
    for r in results:
        p = r["parameters"]
        ws.append([
            r["no"], r["test"], p.get("campus", ""), p.get("city", ""), p.get("vehicle", ""),
            ", ".join(map(str, p.get("excel_rows", []))), str(r["expected"]), str(r["actual"]),
            r["result"], r["failed_step"] or "", r["remark"],
        ])
        cell = ws.cell(row=ws.max_row, column=cols.index("Result") + 1)
        cell.fill = PatternFill("solid", fgColor="C6EFCE" if r["result"] == "PASS" else "FFC7CE")
        cell.font = Font(bold=True)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="305496")
    for idx, width in enumerate([6, 44, 10, 30, 16, 16, 12, 12, 8, 11, 70], start=1):
        ws.column_dimensions[get_column_letter(idx)].width = width
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    sm = wb.create_sheet("Summary", 0)
    for k, v in summary.items():
        sm.append([k.replace("_", " ").title(), v])
    sm.column_dimensions["A"].width = 16
    sm.column_dimensions["B"].width = 70
    for row in sm.iter_rows(min_col=1, max_col=1):
        row[0].font = Font(bold=True)

    xlsx_path = REPORT_DIR / f"cargo_test_report_{stamp}.xlsx"
    wb.save(xlsx_path)
    return json_path, xlsx_path, summary


# ---------------------------------------------------------------- the test
class CargoLoaderTest(TestCase):
    results = []

    @classmethod
    def setUpTestData(cls):
        if not XLSX.exists():
            raise FileNotFoundError(f"Excel sheet not found: {XLSX} (set CARGO_XLSX)")
        Campus = apps.get_model(APP, "Campus")
        Campus.objects.create(pk=CAMPUS_ID, name=CAMPUS_NAME)
        cls.output = StringIO()
        call_command("load_cargo_data", str(XLSX), app=APP, stdout=cls.output)
        cls.groups, cls.student_rows = read_expected(XLSX)

    @classmethod
    def tearDownClass(cls):
        super().tearDownClass()
        if cls.results:
            json_path, xlsx_path, s = write_reports(cls.results)
            print(f"\nCargo test report: {s['passed']}/{s['total']} passed, {s['failed']} failed\n  {json_path}\n  {xlsx_path}")

    def test_dropdown_simulation(self):
        """Pick Campus -> City -> Vehicle for every combination in the sheet and compare the fare."""
        Matrix = apps.get_model(APP, "Matrix")
        campus = apps.get_model(APP, "Campus").objects.get(pk=CAMPUS_ID)
        for (city_key, vehicle_key), g in sorted(self.groups.items(), key=lambda kv: (kv[1]["city"], kv[1]["vehicle"])):
            step, actual, detail = simulate_dropdowns(Matrix, campus, city_key, vehicle_key)
            params = {"campus": campus.name, "city": g["city"], "vehicle": g["vehicle"], "excel_rows": g["rows"]}
            expected = f"{g['expected']:.2f}"
            note = ""
            if g["conflict"]:
                note = ("Sheet has different prices for this city + vehicle "
                        f"({', '.join(f'{p:.2f} x{n}' for p, n in g['conflict'].items())}); most common used. ")
            if step:
                self.results.append(_result(
                    "Dropdown fare lookup", params, expected, None, False, step,
                    f"{note}Step {step} failed. Expected: {g['city']} / {g['vehicle']} = {expected}. Actual: {detail}"))
            elif actual != g["expected"]:
                self.results.append(_result(
                    "Dropdown fare lookup", params, expected, f"{actual:.2f}", False, 4,
                    f"{note}Step 4 fare mismatch. Expected: {expected}. Actual: {actual:.2f}"))
            else:
                self.results.append(_result("Dropdown fare lookup", params, expected, f"{actual:.2f}", True, remark=note.strip()))
        failed = [r for r in self.results if r["result"] == "FAIL"]
        self.assertFalse(failed, f"{len(failed)} dropdown lookups failed - see the report")

    def test_matrix_integrity(self):
        """Whole-table checks: campus, duplicates, extra rows, Campus table, SubsidyType, re-running."""
        Matrix = apps.get_model(APP, "Matrix")
        Campus = apps.get_model(APP, "Campus")
        SubsidyType = apps.get_model(APP, "SubsidyType")
        base = {"campus": CAMPUS_NAME}
        checks = []

        campus_ids = set(Matrix.objects.values_list("campus_id", flat=True))
        checks.append(("Every Matrix row uses campus Main (id 1)", f"{{{CAMPUS_ID}}}", str(campus_ids or "{}"), campus_ids == {CAMPUS_ID}))

        pairs = Counter(Matrix.objects.values_list("city_id", "vehicle_id"))
        dupes = [k for k, n in pairs.items() if n > 1]
        checks.append(("No duplicate Matrix rows per city + vehicle", "0 duplicates", f"{len(dupes)} duplicates", not dupes))

        expected_rows, actual_rows = len(self.groups), Matrix.objects.count()
        checks.append(("Matrix row count equals sheet combinations", str(expected_rows), str(actual_rows), expected_rows == actual_rows))

        checks.append(("No Campus rows created by the loader", "1", str(Campus.objects.count()), Campus.objects.count() == 1))

        has_student = SubsidyType.objects.filter(name="Student").exists()
        checks.append(("SubsidyType 'Student' exists when rebate equals cost",
                       "exists" if self.student_rows else "not needed", "exists" if has_student else "missing",
                       has_student == bool(self.student_rows)))

        before = Matrix.objects.count(), sorted(Matrix.objects.values_list("pk", "peso_price"))
        call_command("load_cargo_data", str(XLSX), app=APP, stdout=StringIO())
        after = Matrix.objects.count(), sorted(Matrix.objects.values_list("pk", "peso_price"))
        checks.append(("Running the loader twice changes nothing", f"{before[0]} rows, same prices",
                       f"{after[0]} rows, {'same' if before[1] == after[1] else 'different'} prices", before == after))

        for name, expected, actual, ok in checks:
            self.results.append(_result(name, dict(base), expected, actual, ok, remark="" if ok else f"Expected: {expected}. Actual: {actual}"))
        failed = [c[0] for c in checks if not c[3]]
        self.assertFalse(failed, f"Integrity checks failed: {failed}")
